#!/usr/bin/env python3
"""Known-answer tests for the Okanagan TPO pipeline core.

Synthetic licences exercise every rule in the locked spec; expected values
are computed independently from first principles (literals below), so a
green run means the arithmetic - not just the plumbing - is right.

Run:  python tests_fixtures.py
"""
from math import isclose

import pandas as pd

from tpo import pipeline as core

SY = core.SECONDS_YEAR  # 31,536,000


def R(lic, pod, purp, qty, units, flag, subtype="POD", conn=None,
      status="Current", holder="Holder A"):
    return dict(LICENCE_NUMBER=lic, POD_NUMBER=pod, PURPOSE_USE_CODE=purp,
                QUANTITY=qty, QUANTITY_UNITS=units, QUANTITY_FLAG=flag,
                POD_SUBTYPE=subtype, HYDRAULIC_CONNECTIVITY=conn,
                LICENCE_STATUS=status, PRIMARY_LICENSEE_NAME=holder)


def fixtures() -> pd.DataFrame:
    rows = [
        # L1 domestic, T, 2 holders; one row carries the trailing-space unit
        R("C100001", "PD001", "01A", 2.27305, "m3/day", "T", holder="A"),
        R("C100001", "PD001", "01A", 2.27305, "m3/day ", "T", holder="B"),
        # L2 irrigation 03B, M across 3 PODs, 2 holders each (6 rows -> 1 comp)
        *[R("C100002", p, "03B", 36500.0, "m3/year", "M", holder=h)
          for p in ("PD010", "PD011", "PD012") for h in ("A", "B")],
        # L3 waterworks, M across two UNITS -> multi-watershed bucket
        R("C100003", "PD020", "00A", 1_000_000.0, "m3/year", "M"),
        R("C100003", "PD021", "00A", 1_000_000.0, "m3/year", "M"),
        # L4 D flag: per-POD known quantities, they sum (same unit)
        R("C100004", "PD030", "02D", 100.0, "m3/day", "D"),
        R("C100004", "PD031", "02D", 200.0, "m3/day", "D"),
        # L5 P flag on wells: connected + unknown buckets
        R("500001", "PW040", "WSA03", 50.0, "m3/day", "P", "PWD", "Likely"),
        R("500001", "PW041", "WSA03", 70.0, "m3/day", "P", "PWD", "Unknown"),
        # L6 null flag on a well, null connectivity -> gw_unknown
        R("500002", "PW050", "WSA08", 1000.0, "m3/year", None, "PWD", None),
        # L7 empty shell (no POD number, no quantity)
        R("C100007", None, "04A", None, None, None),
        # L8 'Total Flow' -> non-quantifiable, counted not summed
        R("C100008", "PD060", "04A", 0.0, "Total Flow", "T"),
        # L9 mixed licence: storage purpose + domestic purpose
        R("C100009", "PD070", "08A", 5000.0, "m3/year", "T"),
        R("C100009", "PD071", "01A", 400.0, "m3/year", "T"),
        # L10 non-current -> filtered
        R("C100010", "PD075", "01A", 999.0, "m3/year", "T", status="Cancelled"),
        # L11 stray multi-POD T (same qty) -> exception, counted once
        R("C100011", "PD080", "02E", 500.0, "m3/year", "T"),
        R("C100011", "PD081", "02E", 500.0, "m3/year", "T"),
        # L12 M with differing quantities -> exception, max used
        R("C100012", "PD090", "00B", 10_000.0, "m3/year", "M"),
        R("C100012", "PD091", "00B", 12_000.0, "m3/year", "M"),
        # L13 POD outside every unit
        R("C100013", "PD100", "01A", 1000.0, "m3/year", "T"),
        # L14 two distinct combos on one licence/POD/purpose (kept separate)
        R("C100014", "PG110", "WSA07", 10.0, "m3/day", None, "PG", None),
        R("C100014", "PG110", "WSA07", 20.0, "m3/day", None, "PG", None),
    ]
    return pd.DataFrame(rows)


POD_UNIT = {"PD001": "mission", "PD010": "mission", "PD011": "mission",
            "PD012": "mission", "PD020": "mission", "PD021": "trout",
            "PD030": "trout", "PD031": "trout", "PW040": "trout",
            "PW041": "trout", "PW050": "trout", "PD060": "mission",
            "PD070": "mission", "PD071": "mission", "PD080": "mission",
            "PD081": "mission", "PD090": "mission", "PD091": "mission",
            "PD100": None, "PG110": "trout"}


def mock_pod_units(raw: pd.DataFrame, cfg: core.Config) -> pd.DataFrame:
    """Replicate the cleaning that assigns ROW_ID, then map PODs to units.
    (This also proves ROW_ID assignment is deterministic across passes,
    which run_pipeline.py relies on.)"""
    led, exc = core.Ledger(), core.Exceptions()
    df = core.normalize_strings(raw)
    df = core.filter_current(df, led)
    df, _ = core.split_shells(df, led)
    df = core.dedup_holders(df, led, exc)
    return df.assign(unit_id=df["POD_NUMBER"].map(POD_UNIT))[["ROW_ID", "unit_id"]]


def get(master, unit, col):
    return float(master.loc[master["unit_id"] == unit, col].iloc[0])


def main() -> None:
    cfg = core.Config("config")
    raw = fixtures()
    tables = core.run_core(raw, mock_pod_units(raw, cfg), cfg)
    master, comps = tables["master"], tables["components"]
    ledger = tables["ledger"].set_index("stage")["rows"]

    # ---- ledger arithmetic -------------------------------------------------
    assert ledger["raw_rows"] == 27
    assert ledger["status_filter_current"] == 26
    assert ledger["shell_records_removed"] == 25          # rows kept
    assert ledger["dedup_six_field_key"] == 21            # 27-1-1-4
    assert ledger["authorization_components"] == 16
    assert ledger["allocation_assigned"] == 14
    assert ledger["allocation_multi_watershed"] == 1
    assert ledger["allocation_outside_units"] == 1
    assert len(tables["shells"]) == 1

    # ---- exceptions --------------------------------------------------------
    kinds = set(tables["exceptions"]["kind"])
    assert {"t_flag_multiple_pods", "m_flag_quantity_mismatch",
            "multiple_combos_same_licence_pod_purpose"} <= kinds

    # ---- hand-computed expectations (independent literals) -----------------
    sw = core.BUCKET_COLUMNS["sw"]
    gwc = core.BUCKET_COLUMNS["gw_connected"]
    gwu = core.BUCKET_COLUMNS["gw_unknown"]

    l1 = 2.27305 * 365                      # 829.66325 m3/yr
    mission_jan = (l1 + 400 + 500 + 12000) / SY
    mission_jul = mission_jan + 36500 * 0.236 / (31 * 86400)
    assert isclose(get(master, "mission", f"{sw}_JAN"), mission_jan, rel_tol=1e-12)
    assert isclose(get(master, "mission", f"{sw}_JUL"), mission_jul, rel_tol=1e-12)
    # irrigation contributes nothing Oct-Mar
    assert isclose(get(master, "mission", f"{sw}_DEC"), mission_jan, rel_tol=1e-12)

    assert isclose(get(master, "trout", f"{sw}_JUN"), (100 + 200) * 365 / SY,
                   rel_tol=1e-12)
    assert isclose(get(master, "trout", f"{gwc}_MAR"), 50 * 365 / SY, rel_tol=1e-12)
    trout_unknown = (70 * 365 + 1000 + (10 + 20) * 365) / SY
    assert isclose(get(master, "trout", f"{gwu}_AUG"), trout_unknown, rel_tol=1e-12)

    # storage identified, not in totals
    assert isclose(get(master, "mission", "STORAGE_IDENTIFIED_ANNUAL_M3"), 5000.0)
    # 'Total Flow' never enters rates
    tf = comps[comps["QUANTITY_UNITS"] == "Total Flow"]
    assert len(tf) == 1 and not bool(tf["quantifiable"].iloc[0])

    # multi-watershed companion: non-additive, mirrored to both touched units
    mwcol = "MULTI_WATERSHED_ASSOCIATED_M3S_NONADDITIVE"
    for u in ("mission", "trout"):
        assert isclose(get(master, u, mwcol), 1_000_000 / SY, rel_tol=1e-12)
    # ...and absent from the summable totals
    assert isclose(get(master, "trout", f"{sw}_JAN"), (100 + 200) * 365 / SY,
                   rel_tol=1e-12)

    # unaffected units stay zero
    assert get(master, "vaseux", f"{sw}_JUL") == 0.0

    # normalization proof: the 'm3/day '/'m3/day' twins collapsed to ONE comp
    assert (comps["LICENCE_NUMBER"] == "C100001").sum() == 1
    # L14's two combos both survived the six-field dedup
    assert (comps["LICENCE_NUMBER"] == "C100014").sum() == 2

    # ---- demo workbook -----------------------------------------------------
    from tpo.workbook import write_workbook
    path = write_workbook("demo_workbook_synthetic.xlsx", tables,
                          extract_note="SYNTHETIC FIXTURE DATA - demo only")
    from openpyxl import load_workbook
    names = load_workbook(path).sheetnames
    assert {"Master", "Monthly_Long", "README", "QA_Ledger",
            "Exceptions"} <= set(names)

    print("ALL CHECKS PASSED")
    print(tables["ledger"].to_string(index=False))
    print("\nMaster (non-zero units, selected cols):")
    cols = ["unit_id", f"{sw}_JAN", f"{sw}_JUL", f"{gwc}_JUL", f"{gwu}_JUL",
            "STORAGE_IDENTIFIED_ANNUAL_M3", mwcol]
    m = master[master["unit_id"].isin(["mission", "trout"])][cols]
    print(m.to_string(index=False))
    print(f"\ndemo workbook: {path} | sheets: {names}")


if __name__ == "__main__":
    main()
