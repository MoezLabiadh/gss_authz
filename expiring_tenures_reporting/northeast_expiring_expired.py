"""
Northeast Region - Land Act tenures EXPIRED (awaiting replacement) and EXPIRING (on or before 2028-03-31)

Steps
  1. BCGW, one query: in-scope tenures with a parcel touching the Northeast NR region (ADM_NR_REGIONS_SP,
     ORG_UNIT = 'RNO'), with managing agency and holders (name, Interested Party ID, address, phone for
     organizations only)
       EXPIRED  = latest expired DTID on a file that has a replacement (REP) application in progress
                  and no tenure in good standing, i.e. expired and awaiting replacement
       EXPIRING = DISPOSITION IN GOOD STANDING with expiry date on/before the cutoff
  2. BCGW: replacement applications in progress, matched to the tenures by file number
  3. BCGW: QA count, by org unit, of in-scope tenures province-wide with no parcel shape (the overlay can't see them)
  4. pandas: one row per DTID
  5. Excel: 'Expired' and 'Expiring' sheets (definitions and assumptions: README.md)

Connection: environment variables BCGW_HOST, BCGW_USER, BCGW_PASSWORD
Requires: oracledb, pandas, openpyxl
"""

import os
import time
from datetime import date, timedelta
from pathlib import Path

import oracledb
import pandas as pd
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

# ------------------------------------------------------------------ settings
NR_REGION = "RNO"                          # ADM_NR_REGIONS_SP.ORG_UNIT for Northeast
EXPIRING_ON_OR_BEFORE = date(2028, 3, 31)
REPLACEMENT_APP_STATUSES = ["ACCEPTED"]  # REP application still in progress
EXCLUDE_TYPES = []                         # e.g. ["RESERVE/NOTATION"]; type counts are printed
BCGW_SERVICE = "idwprod1.bcgov"            # appended when BCGW_HOST is only a host name
OUT_XLSX = Path(f"NE_expired_expiring_tenures_{date.today():%Y%m%d}.xlsx")

REP_STATUSES_SQL = ", ".join(f"'{s}'" for s in REPLACEMENT_APP_STATUSES)


# ------------------------------------------------------------------ SQL
# Every disposition transaction with its current status (used by the queries below)
SQL_TXN = """
    SELECT DT.DISPOSITION_TRANSACTION_SID AS DTID,
           DS.FILE_CHR                    AS FILE_NBR,
           TT.STATUS_NME                  AS STATUS,
           DT.APPLICATION_TYPE_CDE        AS APPLICATION_TYPE,
           DT.EXPIRY_DAT                  AS EXPIRY_DATE
    FROM WHSE_TANTALIS.TA_DISPOSITION_TRANSACTIONS DT
    JOIN WHSE_TANTALIS.TA_DISPOSITIONS DS
      ON DS.DISPOSITION_SID = DT.DISPOSITION_SID
    JOIN WHSE_TANTALIS.TA_DISP_TRANS_STATUSES TS
      ON TS.DISPOSITION_TRANSACTION_SID = DT.DISPOSITION_TRANSACTION_SID
     AND TS.EXPIRY_DAT IS NULL
    JOIN WHSE_TANTALIS.TA_STATUS TT
      ON TT.CODE_CHR = TS.CODE_CHR_STATUS
"""

# Replacement applications in progress, one row per application
SQL_REPLACEMENT_APPS = f"""
WITH txn AS ({SQL_TXN})
SELECT FILE_NBR,
       CAST(DTID AS NUMBER) AS REPLACEMENT_APP_DTID,
       STATUS               AS REPLACEMENT_APP_STATUS
FROM txn
WHERE APPLICATION_TYPE = 'REP'
  AND STATUS IN ({REP_STATUSES_SQL})
"""

# A holder counts as an organization when it has a legal name and no person name. Phone numbers are
# only pulled for organizations (see README, assumptions 9-10).
IS_ORGANIZATION_SQL = "PR.LEGAL_NAME IS NOT NULL AND PR.FIRST_NAME IS NULL AND PR.LAST_NAME IS NULL"

# Does TA_TENANTS say which of the holder's locations (addresses) applies to a tenure?
SQL_TENANT_LOCATION_CHECK = """
SELECT COUNT(*) AS N
FROM all_tab_columns
WHERE owner = 'WHSE_TANTALIS' AND table_name = 'TA_TENANTS' AND column_name = 'LOCATION_SID'
"""

# In-scope tenures, province-wide, one row per DTID: CTEs txn, file_summary and picked
SQL_PICKED_CTES = f"""
txn AS ({SQL_TXN}),
file_summary AS (
    -- one row per Crown land file
    SELECT FILE_NBR,
           MAX(CASE WHEN STATUS = 'EXPIRED' THEN EXPIRY_DATE END)                   AS LAST_EXPIRY_DATE,
           MAX(CASE WHEN STATUS = 'DISPOSITION IN GOOD STANDING' THEN 1 ELSE 0 END) AS HAS_DIGS,
           MAX(CASE WHEN APPLICATION_TYPE = 'REP' AND STATUS IN ({REP_STATUSES_SQL})
                    THEN 1 ELSE 0 END)                                               AS HAS_REPLACEMENT_APP
    FROM txn
    GROUP BY FILE_NBR
),
picked AS (
    -- EXPIRING: in good standing, expiring on/before the cutoff
    -- EXPIRED:  latest expired tenure on a file with a replacement application in progress
    --           and nothing in good standing (= expired, awaiting replacement)
    SELECT t.DTID, t.FILE_NBR, t.STATUS, t.EXPIRY_DATE
    FROM txn t
    JOIN file_summary f
      ON f.FILE_NBR = t.FILE_NBR
    WHERE (t.STATUS = 'DISPOSITION IN GOOD STANDING' AND t.EXPIRY_DATE < :expiring_before)
       OR (t.STATUS = 'EXPIRED' AND t.EXPIRY_DATE = f.LAST_EXPIRY_DATE
           AND f.HAS_REPLACEMENT_APP = 1 AND f.HAS_DIGS = 0)
)
"""

# In-scope tenures with a parcel touching the NE region.
# One row per DTID x interest parcel x tenant. <ADDRESS_FOR_TENURE> is filled in main().
SQL_TENURES = f"""
WITH ne_parcels AS (
    -- parcels touching the NE region. MATERIALIZE runs this spatial join on its own, once, driven by the
    -- spatial index (ORDERED + window table first, as Oracle recommends for SDO_RELATE joins).
    -- If it's slow, ADM_NR_REGIONS_SPG (generalized) is faster, at some cost to boundary precision.
    SELECT /*+ MATERIALIZE ORDERED */ DISTINCT SH.INTRID_SID
    FROM WHSE_ADMIN_BOUNDARIES.ADM_NR_REGIONS_SP R,
         WHSE_TANTALIS.TA_INTEREST_PARCEL_SHAPES SH
    WHERE R.ORG_UNIT = :region
      AND SDO_RELATE(SH.SHAPE, R.SHAPE, 'mask=ANYINTERACT') = 'TRUE'
),
{SQL_PICKED_CTES},
ne_picked AS (
    -- in-scope tenures with a current parcel in the region
    SELECT /*+ MATERIALIZE */ p.*
    FROM picked p
    WHERE p.DTID IN (SELECT IP.DISPOSITION_TRANSACTION_SID
                     FROM ne_parcels NP
                     JOIN WHSE_TANTALIS.TA_INTEREST_PARCELS IP
                       ON IP.INTRID_SID = NP.INTRID_SID
                      AND IP.EXPIRY_DAT IS NULL)
),
holder_address AS (
    -- one address per holder and tenure (README, assumption 8): the one TANTALIS links to the holder
    -- on the tenure, if any, otherwise the most recently updated
    -- (TA_ADDRESSES is keyed by LOCATION_SID + INTERESTED_PARTY_SID + ADDRESS_SID, so the fields are
    -- carried through here rather than joined back on ADDRESS_SID, which is only a sequence number)
    SELECT *
    FROM (SELECT TE.DISPOSITION_TRANSACTION_SID AS DTID, TE.INTERESTED_PARTY_SID,
                 AD.ADDRESS_LINE_1, AD.ADDRESS_LINE_2, AD.ADDRESS_LINE_3,
                 AD.CITY, AD.REGION_CDE, AD.POSTAL_CODE, AD.ZIP_CODE, AD.COUNTRY_CDE,
                 ROW_NUMBER() OVER (
                     PARTITION BY TE.DISPOSITION_TRANSACTION_SID, TE.INTERESTED_PARTY_SID
                     ORDER BY <ADDRESS_FOR_TENURE> DESC,
                              COALESCE(AD.LAST_UPDATE_DAT, AD.CREATE_DAT) DESC NULLS LAST,
                              AD.LOCATION_SID DESC, AD.ADDRESS_SID DESC) AS RN
          FROM ne_picked p
          JOIN WHSE_TANTALIS.TA_TENANTS TE
            ON TE.DISPOSITION_TRANSACTION_SID = p.DTID
          JOIN WHSE_TANTALIS.TA_ADDRESSES AD
            ON AD.INTERESTED_PARTY_SID = TE.INTERESTED_PARTY_SID)
    WHERE RN = 1
)
SELECT CAST(p.DTID AS NUMBER)                  AS DTID,
       CAST(IP.INTRID_SID AS NUMBER)           AS INTRID_SID,
       p.FILE_NBR,
       p.STATUS,
       p.EXPIRY_DATE,
       DT.COMMENCEMENT_DAT                     AS COMMENCEMENT_DATE,
       TY.TYPE_NME                             AS TENURE_TYPE,
       ST.SUBTYPE_NME                          AS TENURE_SUBTYPE,
       PU.PURPOSE_NME                          AS TENURE_PURPOSE,
       SP.SUBPURPOSE_NME                       AS TENURE_SUBPURPOSE,
       DT.DOCUMENT_CHR                         AS DOCUMENT_NBR,
       DT.LOCATION_DSC                         AS LOCATION,
       IP.AREA_HA_NUM                          AS AREA_HA,
       OU.UNIT_NAME                            AS ORG_UNIT,
       DT.MANAGING_AGENCY                      AS MANAGING_AGENCY,
       CAST(TE.INTERESTED_PARTY_SID AS NUMBER) AS INTERESTED_PARTY_ID,
       COALESCE(PR.LEGAL_NAME, TRIM(PR.FIRST_NAME || ' ' || PR.LAST_NAME)) AS CLIENT_NAME,
       CASE WHEN {IS_ORGANIZATION_SQL} THEN 1 ELSE 0 END AS IS_ORGANIZATION,
       CASE WHEN {IS_ORGANIZATION_SQL} THEN PR.WORK_AREA_CODE END        AS ORG_PHONE_AREA_CODE,
       CASE WHEN {IS_ORGANIZATION_SQL} THEN PR.WORK_PHONE_NUMBER END     AS ORG_PHONE_NUMBER,
       CASE WHEN {IS_ORGANIZATION_SQL} THEN PR.WORK_EXTENSION_NUMBER END AS ORG_PHONE_EXTENSION,
       CASE WHEN NOT ({IS_ORGANIZATION_SQL}) AND PR.WORK_PHONE_NUMBER IS NOT NULL
            THEN 1 ELSE 0 END                  AS INDIVIDUAL_HAS_PHONE,  -- QA flag only, no number
       TE.PRIMARY_CONTACT_YRN                  AS PRIMARY_CONTACT,
       TE.SEPARATION_DAT                       AS SEPARATION_DATE,
       HA.ADDRESS_LINE_1, HA.ADDRESS_LINE_2, HA.ADDRESS_LINE_3,
       HA.CITY, HA.REGION_CDE, HA.POSTAL_CODE, HA.ZIP_CODE, HA.COUNTRY_CDE
FROM ne_picked p
JOIN WHSE_TANTALIS.TA_DISPOSITION_TRANSACTIONS DT
  ON DT.DISPOSITION_TRANSACTION_SID = p.DTID
JOIN WHSE_TANTALIS.TA_INTEREST_PARCELS IP
  ON IP.DISPOSITION_TRANSACTION_SID = p.DTID
 AND IP.EXPIRY_DAT IS NULL
-- lookups are LEFT joins so a missing code can't silently drop a tenure
LEFT JOIN WHSE_TANTALIS.TA_AVAILABLE_TYPES TY
  ON TY.TYPE_SID = DT.TYPE_SID
LEFT JOIN WHSE_TANTALIS.TA_AVAILABLE_SUBTYPES ST
  ON ST.SUBTYPE_SID = DT.SUBTYPE_SID AND ST.TYPE_SID = DT.TYPE_SID
LEFT JOIN WHSE_TANTALIS.TA_AVAILABLE_PURPOSES PU
  ON PU.PURPOSE_SID = DT.PURPOSE_SID
LEFT JOIN WHSE_TANTALIS.TA_AVAILABLE_SUBPURPOSES SP
  ON SP.SUBPURPOSE_SID = DT.SUBPURPOSE_SID AND SP.PURPOSE_SID = DT.PURPOSE_SID
LEFT JOIN WHSE_TANTALIS.TA_ORGANIZATION_UNITS OU
  ON OU.ORG_UNIT_SID = DT.ORG_UNIT_SID
-- all tenants, separated or not: pandas keeps the current ones (or the last ones if none are current)
LEFT JOIN WHSE_TANTALIS.TA_TENANTS TE
  ON TE.DISPOSITION_TRANSACTION_SID = p.DTID
LEFT JOIN WHSE_TANTALIS.TA_INTERESTED_PARTIES PR
  ON PR.INTERESTED_PARTY_SID = TE.INTERESTED_PARTY_SID
LEFT JOIN holder_address HA
  ON HA.DTID = p.DTID AND HA.INTERESTED_PARTY_SID = TE.INTERESTED_PARTY_SID
"""

# QA: in-scope tenures, province-wide, with a current parcel but no parcel shape, so the overlay can't see them
SQL_NO_SHAPE = f"""
WITH {SQL_PICKED_CTES}
SELECT OU.UNIT_NAME AS ORG_UNIT, COUNT(DISTINCT p.DTID) AS TENURES
FROM picked p
JOIN WHSE_TANTALIS.TA_DISPOSITION_TRANSACTIONS DT
  ON DT.DISPOSITION_TRANSACTION_SID = p.DTID
LEFT JOIN WHSE_TANTALIS.TA_ORGANIZATION_UNITS OU
  ON OU.ORG_UNIT_SID = DT.ORG_UNIT_SID
WHERE EXISTS (SELECT 1 FROM WHSE_TANTALIS.TA_INTEREST_PARCELS IP
              WHERE IP.DISPOSITION_TRANSACTION_SID = p.DTID AND IP.EXPIRY_DAT IS NULL)
  AND NOT EXISTS (SELECT 1
                  FROM WHSE_TANTALIS.TA_INTEREST_PARCELS IP
                  JOIN WHSE_TANTALIS.TA_INTEREST_PARCEL_SHAPES SH
                    ON SH.INTRID_SID = IP.INTRID_SID
                  WHERE IP.DISPOSITION_TRANSACTION_SID = p.DTID AND IP.EXPIRY_DAT IS NULL)
GROUP BY OU.UNIT_NAME
ORDER BY TENURES DESC
"""

TENURE_COLS = [
    "FILE_NBR", "STATUS", "EXPIRY_DATE", "COMMENCEMENT_DATE", "TENURE_TYPE", "TENURE_SUBTYPE",
    "TENURE_PURPOSE", "TENURE_SUBPURPOSE", "DOCUMENT_NBR", "LOCATION", "ORG_UNIT", "MANAGING_AGENCY",
]

# output columns, per sheet
DETAIL_COLS = [
    "NR_REGION", "ORG_UNIT", "MANAGING_AGENCY", "CLIENT_NAME", "INTERESTED_PARTY_ID", "CLIENT_ADDRESS",
    "CLIENT_PHONE", "CLIENT_STATUS",
    "TENURE_TYPE", "TENURE_SUBTYPE", "TENURE_PURPOSE", "TENURE_SUBPURPOSE", "DOCUMENT_NBR", "COMMENCEMENT_DATE",
    "AREA_HA", "LOCATION", "INTEREST_PARCEL_IDS",
]
EXPIRED_COLS = (["FILE_NBR", "DTID", "REPLACEMENT_APP_DTID", "REPLACEMENT_APP_STATUS", "STATUS", "EXPIRY_DATE"]
                + DETAIL_COLS)
EXPIRING_COLS = (["FILE_NBR", "DTID", "STATUS", "EXPIRY_DATE", "PAST_EXPIRY_DATE"] + DETAIL_COLS
                 + ["REPLACEMENT_APP_DTID", "REPLACEMENT_APP_STATUS"])


# ------------------------------------------------------------------ functions
def run_query(conn, sql, params=None):
    with conn.cursor() as cur:
        cur.arraysize = 5000
        cur.execute(sql, params)
        cols = [d[0] for d in cur.description]
        return pd.DataFrame(cur.fetchall(), columns=cols)


def txt(v):
    """Stripped string, or None when blank."""
    return v.strip() if isinstance(v, str) and v.strip() else None


def format_address(r):
    """One-line address from TA_ADDRESSES fields, e.g. 'PO BOX 1, UNIT 2, FORT ST. JOHN BC V1J 4H6, CA'."""
    city_line = " ".join(p for p in [txt(r.CITY), txt(r.REGION_CDE), txt(r.POSTAL_CODE) or txt(r.ZIP_CODE)] if p)
    parts = [txt(r.ADDRESS_LINE_1), txt(r.ADDRESS_LINE_2), txt(r.ADDRESS_LINE_3), city_line, txt(r.COUNTRY_CDE)]
    return ", ".join(p for p in parts if p) or None


def format_phone(r):
    """Work phone from TA_INTERESTED_PARTIES fields, e.g. '(250) 555-1234 ext. 12'.
    None without a number (an area code alone is not a phone number)."""
    number = txt(r.ORG_PHONE_NUMBER)
    if not number:
        return None
    if len(number) == 7 and number.isdigit():
        number = f"{number[:3]}-{number[3:]}"
    area, ext = txt(r.ORG_PHONE_AREA_CODE), txt(r.ORG_PHONE_EXTENSION)
    phone = f"({area}) {number}" if area else number
    return f"{phone} ext. {ext}" if ext else phone


def build_report(rows, rep_apps):
    """Returns (expired, expiring) DataFrames, one row per tenure (DTID)."""
    today = pd.Timestamp.today().normalize()
    rows = rows.copy()
    rows["DTID"] = rows["DTID"].astype("int64")
    rows["INTRID_SID"] = rows["INTRID_SID"].astype("int64")
    for c in ["EXPIRY_DATE", "COMMENCEMENT_DATE", "SEPARATION_DATE"]:
        rows[c] = pd.to_datetime(rows[c], errors="coerce")
        rows[c] = rows[c].where(rows[c].dt.year.between(1900, 2999))  # legacy typos (e.g. year 0201) -> blank

    # --- one row per tenure (rows are DTID x parcel x tenant)
    parcels = rows.drop_duplicates(["DTID", "INTRID_SID"])
    ten = rows.drop_duplicates("DTID").set_index("DTID")[TENURE_COLS].copy()
    ten["INTEREST_PARCEL_IDS"] = parcels.groupby("DTID")["INTRID_SID"].agg(
        lambda s: "; ".join(map(str, sorted(s))))
    ten["AREA_HA"] = parcels.groupby("DTID")["AREA_HA"].sum(min_count=1)

    # --- clients: current tenants (no separation date), primary contact first.
    #     A tenure with no current tenant shows its last tenant(s) instead.
    #     The SQL already picked one address per holder and tenure.
    holders = rows.dropna(subset=["INTERESTED_PARTY_ID"])
    t = (holders.sort_values("SEPARATION_DATE", ascending=False, na_position="first")
                .drop_duplicates(["DTID", "INTERESTED_PARTY_ID"]))
    t = t.assign(CURRENT=t["SEPARATION_DATE"].isna(),
                 CLIENT_ADDRESS=[format_address(r) for r in t.itertuples(index=False)])
    has_current = t.groupby("DTID")["CURRENT"].transform("any")
    last_sep = t.groupby("DTID")["SEPARATION_DATE"].transform("max")
    t = t[t["CURRENT"] | (~has_current & t["SEPARATION_DATE"].eq(last_sep))]
    t = t.sort_values(["DTID", "PRIMARY_CONTACT"], ascending=[True, False])

    # phone: organizations only (SQL returns no number for individuals)
    org = t["IS_ORGANIZATION"].eq(1)
    phone = pd.Series([format_phone(r) for r in t.itertuples(index=False)], index=t.index, dtype="string")
    t["CLIENT_PHONE"] = phone.fillna("(no phone)").where(org, "(individual)")

    print(f"\nHolders: {len(t):,} ({org.sum():,} organizations, {(~org).sum():,} individuals)")
    print(f"  no address on file: {t['CLIENT_ADDRESS'].isna().sum():,}")
    print(f"  organizations with no phone on file: {(org & phone.isna()).sum():,}")
    print(f"  individuals with a work phone recorded in TANTALIS (not included): "
          f"{(~org & t['INDIVIDUAL_HAS_PHONE'].eq(1)).sum():,}")
    clients = t.groupby("DTID").agg(
        CLIENT_NAME=("CLIENT_NAME", lambda s: "; ".join(s.fillna("?"))),
        INTERESTED_PARTY_ID=("INTERESTED_PARTY_ID", lambda s: "; ".join(str(int(v)) for v in s)),
        CLIENT_ADDRESS=("CLIENT_ADDRESS", lambda s: "; ".join(s.fillna("(no address)"))),
        CLIENT_PHONE=("CLIENT_PHONE", lambda s: "; ".join(s)),
        CLIENT_STATUS=("SEPARATION_DATE",
                       lambda s: "CURRENT" if s.isna().all() else f"SEPARATED {s.max():%Y-%m-%d}"),
    )
    ten = ten.join(clients)

    # --- replacement application(s) in progress on the same file, in matching order
    rep = (rep_apps.sort_values("REPLACEMENT_APP_DTID")
                   .groupby("FILE_NBR")
                   .agg(REPLACEMENT_APP_DTID=("REPLACEMENT_APP_DTID", lambda s: "; ".join(str(int(v)) for v in s)),
                        REPLACEMENT_APP_STATUS=("REPLACEMENT_APP_STATUS", "; ".join)))
    ten = ten.join(rep, on="FILE_NBR")

    ten["SHEET"] = ten["STATUS"].map({"EXPIRED": "Expired"}).fillna("Expiring")
    past = (ten["SHEET"] == "Expiring") & (ten["EXPIRY_DATE"] < today)
    ten["PAST_EXPIRY_DATE"] = past.map({True: "Y", False: "N"})
    ten["NR_REGION"] = NR_REGION

    # --- console summary (types before EXCLUDE_TYPES, so you can see what there is to drop)
    print("\nBy tenure type (set EXCLUDE_TYPES to drop any):")
    print(pd.crosstab(ten["TENURE_TYPE"].fillna("(blank)"), ten["SHEET"]).to_string())
    ten = ten[~ten["TENURE_TYPE"].isin(EXCLUDE_TYPES)].copy()

    summary = ten.assign(CLIENT_STATUS=ten["CLIENT_STATUS"].fillna("NO TENANT RECORD").astype(str)
                         .str.replace(r"^SEPARATED.*", "SEPARATED", regex=True))
    for col in ["ORG_UNIT", "MANAGING_AGENCY", "REPLACEMENT_APP_STATUS", "CLIENT_STATUS"]:
        print(f"\nBy {col.lower().replace('_', ' ')}:")
        print(pd.crosstab(summary[col].fillna("(blank)"), summary["SHEET"]).to_string())
    print(f"\nExpired, awaiting replacement: {(ten['SHEET'] == 'Expired').sum():,}")
    print(f"Expiring on or before {EXPIRING_ON_OR_BEFORE}: {(ten['SHEET'] == 'Expiring').sum():,} "
          f"(expiry date already passed: {(ten['PAST_EXPIRY_DATE'] == 'Y').sum():,})")

    ten = ten.reset_index().sort_values(["EXPIRY_DATE", "FILE_NBR"])
    for c in ["EXPIRY_DATE", "COMMENCEMENT_DATE"]:
        ten[c] = ten[c].dt.date
    return (ten.loc[ten["SHEET"] == "Expired", EXPIRED_COLS],
            ten.loc[ten["SHEET"] == "Expiring", EXPIRING_COLS])


def write_excel(sheets, path):
    """One Excel table per sheet, style 'Blue, Table Style Medium 9'."""
    with pd.ExcelWriter(path, engine="openpyxl") as xw:
        for name, df in sheets.items():
            df.to_excel(xw, sheet_name=name, index=False)
            ws = xw.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.style = "Normal"  # drop pandas' bold header so the table style shows
            # a table needs at least one row under the header
            table = Table(displayName=name, ref=f"A1:{get_column_letter(df.shape[1])}{max(len(df) + 1, 2)}")
            table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
            ws.add_table(table)
            for col in ws.columns:
                longest = max(len(str(c.value)) if c.value is not None else 0 for c in col[:500])
                ws.column_dimensions[col[0].column_letter].width = min(max(longest + 4, 10), 50)  # +4: filter button


def main():
    host = os.environ["BCGW_HOST"]
    dsn = host if "/" in host else f"{host}/{BCGW_SERVICE}"  # accepts "host" or "host/service"

    expiring_before = {"expiring_before": EXPIRING_ON_OR_BEFORE + timedelta(days=1)}
    with oracledb.connect(user=os.environ["BCGW_USER"], password=os.environ["BCGW_PASSWORD"], dsn=dsn) as conn:
        # prefer the address TANTALIS links to the holder on each tenure, if this schema has that link
        has_location = run_query(conn, SQL_TENANT_LOCATION_CHECK)["N"].iloc[0] > 0
        print("Addresses: " + ("the holder's location on each tenure (TA_TENANTS.LOCATION_SID)" if has_location
                               else "no TA_TENANTS.LOCATION_SID, so each holder's most recently updated address"))
        sql = SQL_TENURES.replace("<ADDRESS_FOR_TENURE>",
                                  "CASE WHEN AD.LOCATION_SID = TE.LOCATION_SID THEN 1 ELSE 0 END"
                                  if has_location else "0")

        print(f"Expired (awaiting replacement) and expiring tenures touching {NR_REGION}...")
        start = time.perf_counter()
        rows = run_query(conn, sql, {**expiring_before, "region": NR_REGION})
        print(f"  {rows['DTID'].nunique():,} tenures, {len(rows):,} rows ({time.perf_counter() - start:.0f} s)")
        if rows.empty:
            raise SystemExit(f"No expired/expiring tenures touch region '{NR_REGION}' - check NR_REGION")
        rep_apps = run_query(conn, SQL_REPLACEMENT_APPS)

        print("QA: in-scope tenures with no parcel shape, province-wide...")
        no_shape = run_query(conn, SQL_NO_SHAPE, expiring_before)
        if len(no_shape):
            print(f"  {no_shape['TENURES'].sum():,} tenures can't be seen by the overlay. By org unit:")
            print(no_shape.set_index("ORG_UNIT")["TENURES"].to_string())

    expired, expiring = build_report(rows, rep_apps)
    write_excel({"Expired": expired, "Expiring": expiring}, OUT_XLSX)
    print(f"\nWrote {len(expired):,} expired and {len(expiring):,} expiring tenures to {OUT_XLSX}")


if __name__ == "__main__":
    main()
